"""Excel (.xlsx) reading shared by every spreadsheet import: sheet inspection,
header-row guessing, merged cells, blanks, dates, never-calculated formulas,
caps and refusals — plus the /api/excel endpoints and .xlsx datasets, which
must come out exactly as the same data uploaded as CSV.

Workbooks are built in code with openpyxl (no binary fixtures).
"""

import io
import zipfile
from datetime import datetime

import mongomock
import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from backend.services import excel

A = "user-a"
USERS = {A: {"uid": A, "email": "a@x.com", "name": "A"}}


def _xlsx(build) -> bytes:
    wb = Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmea_sheet(wb):
    ws = wb.active
    ws.title = "FMEA"
    ws["A1"] = "Pump FMEA — issue 3"
    ws.merge_cells("A1:F1")
    ws.append([])
    ws.append(["Function", "Functional failure", "Failure mode", "Effect", None, "Date"])
    ws.merge_cells("D3:E3")  # a header spanning two columns
    ws.append(["Pump water", "No flow", "Bearing seized", "Stops", "trip", datetime(2024, 3, 14)])
    ws.append([None, None, "Impeller worn", "Low flow", "alarm", datetime(2024, 3, 14, 8, 30)])
    ws.merge_cells("A4:A5")  # the function cell spans both modes
    ws.append([])
    ws.append(["", "", "Seal leak", None, None, 5.0])
    ws["D7"] = "=D5&\"!\""  # a formula never calculated (openpyxl saves no value)
    data = wb.create_sheet("Data")
    data.append(["hours", "failed"])
    data.append([1240.5, 1])
    data.append([980, 0])


# ---------------------------------------------------------------------------
# Service
# ---------------------------------------------------------------------------


def test_inspect_lists_sheets_with_preview_and_guessed_header():
    info = excel.inspect(_xlsx(_fmea_sheet), "pump.xlsx")
    fmea, data = info["sheets"]
    assert [fmea["name"], data["name"]] == ["FMEA", "Data"]
    # The one-cell title (merged across the table) and the blank row aren't the header.
    assert fmea["guessed_header_row"] == 3
    assert data["guessed_header_row"] == 1
    assert fmea["rows"] == 5 and fmea["cols"] == 6
    assert fmea["row_numbers"][:3] == [1, 2, 3]       # blank rows kept, so numbers match Excel
    assert fmea["preview"][0][0] == "Pump FMEA — issue 3"
    assert data["preview"][1] == ["1240.5", "1"]
    assert any("no saved result" in n for n in fmea["notes"])
    assert data["notes"] == []


def test_read_table_handles_merges_blanks_dates_and_formulas():
    table = excel.read_table(_xlsx(_fmea_sheet), "FMEA", 3)
    assert table.header == ["Function", "Functional failure", "Failure mode", "Effect", "Effect (2)", "Date"]
    assert table.row_numbers == [4, 5, 7]              # blank row 6 skipped
    first, second, third = table.rows
    assert first == ["Pump water", "No flow", "Bearing seized", "Stops", "trip", "2024-03-14"]
    # The vertically merged function cell fills down; the plain blank below it doesn't.
    assert second[0] == "Pump water" and second[1] is None
    assert second[5] == "2024-03-14 08:30:00"
    assert third[0] is None and third[2] == "Seal leak"
    assert third[3] is None                              # uncached formula reads blank …
    assert third[5] == 5                                 # … and 5.0 stays a whole number
    assert any("no saved result" in n for n in table.notes)


def test_read_table_guesses_the_header_when_not_given():
    table = excel.read_table(_xlsx(_fmea_sheet), "FMEA", None)
    assert table.header_row == 3 and table.header[0] == "Function"


def test_no_header_row_names_columns_like_a_headerless_csv():
    table = excel.read_table(_xlsx(_fmea_sheet), "Data", 0)
    assert table.header == ["col 1", "col 2"]
    assert table.rows[0] == ["hours", "failed"]


def test_blank_and_duplicate_headers_are_made_unique():
    def build(wb):
        ws = wb.active
        ws.append(["time", None, "time", "censor"])
        ws.append([1, 2, 3, 0])
    table = excel.read_table(_xlsx(build))
    assert table.header == ["time", "Column B", "time (2)", "censor"]


def test_error_cells_read_blank_with_a_note():
    def build(wb):
        ws = wb.active
        ws.append(["x", "y"])
        ws.append([1, "#DIV/0!"])
        ws["B2"].data_type = "e"
    table = excel.read_table(_xlsx(build))
    assert table.rows[0] == [1, None]
    assert any("Excel error" in n for n in table.notes)


def test_to_csv_and_dataframe_match_the_equivalent_csv():
    table = excel.read_table(_xlsx(_fmea_sheet), "Data")
    assert excel.to_csv(table) == b"hours,failed\n1240.5,1\n980,0\n"
    df = excel.to_dataframe(table)
    pd.testing.assert_frame_equal(df, pd.read_csv(io.StringIO("hours,failed\n1240.5,1\n980,0\n")))


def test_empty_sheet_and_unknown_sheet_are_refused():
    def build(wb):
        wb.active.title = "Empty"
    data = _xlsx(build)
    with pytest.raises(excel.ExcelError, match="no data"):
        excel.read_table(data, "Empty")
    with pytest.raises(excel.ExcelError, match="no sheet named"):
        excel.read_table(data, "Nope")


def test_row_cap_notes_truncation(monkeypatch):
    def build(wb):
        ws = wb.active
        ws.append(["t"])
        for i in range(50):
            ws.append([i])
    data = _xlsx(build)
    table = excel.read_table(data, max_rows=10)
    assert len(table.rows) == 10
    assert any("first 10 rows" in n for n in table.notes)
    monkeypatch.setattr(excel, "MAX_ROWS", 20)
    assert excel.inspect(data)["sheets"][0]["rows"] == 20


def test_column_cap(monkeypatch):
    monkeypatch.setattr(excel, "MAX_COLS", 3)
    def build(wb):
        ws = wb.active
        ws.append(["a", "b", "c", "d", "e"])
        ws.append([1, 2, 3, 4, 5])
    table = excel.read_table(_xlsx(build))
    assert table.header == ["a", "b", "c"]
    assert any("first 3 columns" in n for n in table.notes)


def test_refuses_non_xlsx_legacy_encrypted_and_oversized(monkeypatch):
    with pytest.raises(excel.ExcelError, match="empty"):
        excel.inspect(b"", "a.xlsx")
    with pytest.raises(excel.ExcelError, match="isn't an Excel"):
        excel.inspect(b"time,censor\n1,0\n", "a.xlsx")
    cfb = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512
    with pytest.raises(excel.ExcelError, match="legacy Excel 97"):
        excel.inspect(cfb, "old.xls")
    encrypted = cfb + "EncryptionInfo".encode("utf-16-le") + b"\0" * 64
    with pytest.raises(excel.ExcelError, match="password-protected"):
        excel.inspect(encrypted, "secret.xlsx")
    with pytest.raises(excel.ExcelError, match="password-protected"):
        excel.inspect(cfb, "secret.xlsx")      # an OLE file named .xlsx is an encrypted workbook
    with pytest.raises(excel.ExcelError, match=r"\.xlsb"):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/workbook.bin", b"\0")
        excel.inspect(buf.getvalue(), "book.xlsb")
    with pytest.raises(excel.ExcelError, match=r"OpenDocument"):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            zf.writestr("content.xml", "<x/>")
        excel.inspect(buf.getvalue(), "book.ods")
    with pytest.raises(excel.ExcelError, match="no workbook inside"):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("hello.txt", "hi")
        excel.inspect(buf.getvalue(), "a.xlsx")
    with pytest.raises(excel.ExcelError, match="damaged"):
        excel.inspect(b"PK\x03\x04" + b"\0" * 100, "a.xlsx")

    data = _xlsx(_fmea_sheet)
    monkeypatch.setattr(excel, "MAX_FILE_BYTES", len(data) - 1)
    with pytest.raises(excel.ExcelError, match="larger than"):
        excel.inspect(data, "a.xlsx")


def test_zip_bomb_guard(monkeypatch):
    data = _xlsx(_fmea_sheet)
    monkeypatch.setattr(excel, "MAX_UNCOMPRESSED_BYTES", 1000)
    with pytest.raises(excel.ExcelError, match="expands to more data"):
        excel.inspect(data, "a.xlsx")


def test_xml_is_parsed_with_defusedxml():
    """openpyxl switches to defusedxml when it's installed; the merge/formula
    pass uses defusedxml directly. An entity-expansion payload is refused
    rather than expanded."""
    from openpyxl.xml import DEFUSEDXML

    assert DEFUSEDXML is True
    data = _xlsx(_fmea_sheet)
    src, out = io.BytesIO(data), io.BytesIO()
    bomb = (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa">'
            b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>')
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if item.filename == "xl/sharedStrings.xml" or item.filename.endswith("sheet1.xml"):
                body = bomb + (body.split(b"?>", 1)[1] if body.startswith(b"<?xml") else body)
            zout.writestr(item, body)
    with pytest.raises(excel.ExcelError):
        excel.inspect(out.getvalue(), "a.xlsx")


def test_is_excel_by_extension_or_content():
    assert excel.is_excel(b"PK\x03\x04...", "data")
    assert excel.is_excel(b"x", "Book.XLSX")
    assert excel.is_excel(b"x", "old.xls")
    assert not excel.is_excel(b"time,censor\n1,0\n", "data.csv")


def test_guess_mapping_prefers_exact_then_longest_match():
    fields = [
        {"id": "mode", "label": "Failure mode", "synonyms": ["cause"]},
        {"id": "failure", "label": "Functional failure", "synonyms": ["failure"]},
        {"id": "effects", "label": "Effects", "synonyms": ["effect"]},
    ]
    header = ["Cause", "Failure Mode (description)", "Functional failure", "Local effect"]
    # "Cause" is an exact synonym, so it beats the header that merely contains
    # "failure mode"; each column is used once.
    assert excel.guess_mapping(header, fields) == {
        "mode": "Cause", "failure": "Functional failure", "effects": "Local effect",
    }


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------


@pytest.fixture()
def client(monkeypatch):
    from fastapi.testclient import TestClient

    from backend import config, db
    from backend.auth import get_current_user
    from backend.main import app

    monkeypatch.setattr(config, "AUTH_DISABLED", False)
    monkeypatch.setattr(config, "BILLING_ENABLED", False)
    test_db = mongomock.MongoClient()["reliafy_test"]
    monkeypatch.setattr(db, "_db", test_db)
    monkeypatch.setattr(db, "_simulated", True)
    app.dependency_overrides[get_current_user] = lambda: USERS[A]
    tc = TestClient(app)
    tc.db = test_db
    try:
        yield tc
    finally:
        app.dependency_overrides.clear()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _files(data, name="pump.xlsx"):
    return {"file": (name, data, XLSX)}


def test_inspect_endpoint(client):
    r = client.post("/api/excel/inspect", files=_files(_xlsx(_fmea_sheet)))
    assert r.status_code == 200, r.text
    assert [s["name"] for s in r.json()["sheets"]] == ["FMEA", "Data"]


def test_endpoints_refuse_bad_files_with_422(client):
    for path in ("/api/excel/inspect", "/api/excel/table", "/api/excel/csv"):
        r = client.post(path, files=_files(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "old.xls"))
        assert r.status_code == 422 and "save it as .xlsx" in r.json()["detail"].lower()


def test_endpoints_never_500_on_unexpected_errors(client, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("parser exploded")
    monkeypatch.setattr(excel, "inspect", boom)
    r = client.post("/api/excel/inspect", files=_files(_xlsx(_fmea_sheet)))
    assert r.status_code == 422 and "damaged" in r.json()["detail"]


def test_table_endpoint_with_target_returns_fields_and_guess(client):
    r = client.post("/api/excel/table", files=_files(_xlsx(_fmea_sheet)),
                    data={"sheet": "FMEA", "header_row": "3", "target": "rcm"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["columns"][:3] == ["Function", "Functional failure", "Failure mode"]
    assert body["n_rows"] == 3 and body["rows"][0][5] == "2024-03-14"
    assert body["guess"]["mode"] == "Failure mode" and body["guess"]["function"] == "Function"
    assert any(f["id"] == "mode" and f.get("required") for f in body["fields"])
    plain = client.post("/api/excel/table", files=_files(_xlsx(_fmea_sheet)), data={"sheet": "Data"}).json()
    assert "fields" not in plain and plain["columns"] == ["hours", "failed"]


def test_csv_endpoint_converts_one_sheet(client):
    r = client.post("/api/excel/csv", files=_files(_xlsx(_fmea_sheet)), data={"sheet": "Data", "header_row": "1"})
    assert r.status_code == 200, r.text
    assert r.text == "hours,failed\n1240.5,1\n980,0\n"
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["x-excel-rows"] == "2"
    assert 'filename="pump.csv"' in r.headers["content-disposition"]


def test_excel_endpoints_require_auth(client):
    from backend.main import app

    app.dependency_overrides.clear()
    r = client.post("/api/excel/inspect", files=_files(_xlsx(_fmea_sheet)))
    assert r.status_code == 401


# ---------------------------------------------------------------------------
# Datasets: .xlsx == the same data as CSV
# ---------------------------------------------------------------------------


def _life_frame():
    rng = np.random.default_rng(3)
    return pd.DataFrame({
        "hours": np.round(rng.weibull(1.8, 30) * 1000, 2),
        "censored": rng.integers(0, 2, 30),
        "batch": [f"B{i % 3}" for i in range(30)],
    })


def _life_xlsx(df, title_rows=0):
    def build(wb):
        ws = wb.active
        ws.title = "Bearings"
        for _ in range(title_rows):
            ws.append(["Bearing life test"])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([v.item() if hasattr(v, "item") else v for v in row])
        other = wb.create_sheet("Notes")
        other.append(["nothing to see"])
    return _xlsx(build)


def test_dataset_from_xlsx_matches_the_csv_upload(client):
    df = _life_frame()
    csv_bytes = df.to_csv(index=False).encode()
    via_csv = client.post("/api/datasets", files={"file": ("bearings.csv", csv_bytes, "text/csv")}).json()
    via_xlsx = client.post(
        "/api/datasets", files=_files(_life_xlsx(df, title_rows=2), "bearings.xlsx"),
        data={"sheet": "Bearings", "header_row": "3"},
    )
    assert via_xlsx.status_code == 200, via_xlsx.text
    body = via_xlsx.json()
    assert body["n_rows"] == via_csv["n_rows"] == 30
    assert body["columns"] == via_csv["columns"]
    assert body["preview"] == via_csv["preview"]

    fit_csv = client.post("/api/fit/weibull", data={"dataset_id": via_csv["id"], "x": "hours", "c": "censored"}).json()
    fit_xlsx = client.post("/api/fit/weibull", data={"dataset_id": body["id"], "x": "hours", "c": "censored"}).json()
    assert fit_xlsx["params"] == fit_csv["params"]


def test_dataset_from_xlsx_guesses_header_and_takes_first_sheet(client):
    df = _life_frame()
    r = client.post("/api/datasets", files=_files(_life_xlsx(df, title_rows=1), "bearings.xlsx"))
    assert r.status_code == 200, r.text
    assert [c["name"] for c in r.json()["columns"]] == ["hours", "censored", "batch"]
    assert r.json()["name"] == "bearings"  # the file name, without its extension


def test_dataset_from_xlsx_with_no_header(client):
    df = _life_frame()
    r = client.post("/api/datasets", files=_files(_life_xlsx(df), "b.xlsx"), data={"no_header": "true"})
    assert r.status_code == 200, r.text
    assert [c["name"] for c in r.json()["columns"]] == ["col 1", "col 2", "col 3"]
    assert r.json()["n_rows"] == 31


def test_dataset_refuses_legacy_xls(client):
    r = client.post("/api/datasets", files={"file": ("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "application/vnd.ms-excel")})
    assert r.status_code == 422 and ".xlsx" in r.json()["detail"]


def test_csv_dataset_upload_is_unchanged(client):
    r = client.post("/api/datasets", files={"file": ("d.csv", b"t,c\n1,0\n2,1\n", "text/csv")}, data={"no_header": "false"})
    assert r.status_code == 200 and r.json()["n_rows"] == 2
    assert [c["name"] for c in r.json()["columns"]] == ["t", "c"]
