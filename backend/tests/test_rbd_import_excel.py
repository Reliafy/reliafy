"""RBD import from Excel: the downloadable template (which must itself import
and analyse), two-sheet and single-sheet layouts, block types, parameter
columns, repair data, errors, the column-mapping fallback and the API."""

import io
import json

import mongomock
import pytest
from openpyxl import Workbook, load_workbook

from backend.services import rbd_import
from backend.services.rbd_graph import normalize_graph
from backend.services.rbd_import import excel as rx
from backend.services.rbd_import import RbdImportError


def _book(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _one(data, name="plant.xlsx", mapping=None):
    (d,) = rbd_import.import_file(data, name, mapping)
    return d


def _nodes(graph):
    return {n["id"]: n for n in graph["nodes"]}


def _edges(graph):
    return {(e["source"], e["target"]) for e in graph["edges"]}


# ---------------------------------------------------------------------------
# The template
# ---------------------------------------------------------------------------


def test_template_has_readme_blocks_and_connections():
    wb = load_workbook(io.BytesIO(rx.build_template()), read_only=True)
    assert wb.sheetnames == ["README", "Blocks", "Connections"]
    readme = " ".join(str(c) for row in wb["README"].iter_rows(values_only=True) for c in row if c)
    for word in ("Weibull", "Exponential", "k-of-n", "standby", "MTTR", "Feeds"):
        assert word in readme


def test_template_imports_and_analyses():
    import matplotlib

    matplotlib.use("Agg")
    from backend.services import rbds as rbds_service

    d = _one(rx.build_template(), "reliafy-rbd-template.xlsx")
    assert d.source_format == "Excel" and d.name == "reliafy-rbd-template" and d.warnings == []
    g = d.graph
    nodes = _nodes(g)
    assert nodes["P1"]["model"] == {"distribution_id": "weibull",
                                    "params": [{"name": "alpha", "value": 5000.0}, {"name": "beta", "value": 1.4}]}
    assert nodes["V"]["type"] == "knode" and nodes["V"]["n"] == 2 and nodes["V"]["k"] == 3
    assert nodes["C"]["type"] == "standby" and nodes["C"]["spares"] == 1 and nodes["C"]["cold"] is True
    assert nodes["C"]["model"]["params"] == [{"name": "failure_rate", "value": 1 / 20000}]
    assert nodes["OV"]["model"]["distribution_id"] == "lognormal"
    assert ("input", "F1") in _edges(g) and ("OV", "output") in _edges(g)
    assert g["unit"] == "Hours" and not g.get("repairable")

    graph = normalize_graph(g)
    db = mongomock.MongoClient()["t"]
    assert rbds_service.validate_graph(db, graph, ["u"])["valid"]
    result = rbds_service.analyze_graph(db, graph, ["u"])
    assert 0 < result["mttf"] < 8000


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------


def test_two_sheets_with_other_names_and_synonyms():
    data = _book({
        "Instructions": [["read me"]],
        "Components": [
            ["Tag", "Component", "Dist", "Eta", "Shape", "MTBF"],
            ["A", "Motor", "Weibull", 1000, 2, None],
            ["B", "Gearbox", "Exponential", None, None, 4000],
            ["C", "Belt", None, None, None, 500],          # blank distribution + MTBF -> exponential
        ],
        "Links": [["Source", "Target"], ["Motor", "B"], ["B", "C"]],
    })
    d = _one(data)
    nodes = _nodes(d.graph)
    assert nodes["A"]["label"] == "Motor"
    assert nodes["A"]["model"]["params"] == [{"name": "alpha", "value": 1000.0}, {"name": "beta", "value": 2.0}]
    assert nodes["C"]["model"] == {"distribution_id": "exponential", "params": [{"name": "failure_rate", "value": 0.002}]}
    assert _edges(d.graph) == {("input", "A"), ("A", "B"), ("B", "C"), ("C", "output")}


def test_single_sheet_with_feeds_column_and_block_types():
    data = _book({"Plant": [
        ["Name", "Type", "Distribution", "Failure rate", "Count", "Dormancy", "Feeds"],
        ["Supply", "", "Exponential", 1e-4, None, None, "Pump 1, Pump 2; Pump 3"],
        ["Pump 1", "", "Exponential", 1e-3, None, None, "Vote"],
        ["Pump 2", "", "Exponential", 1e-3, None, None, "Vote"],
        ["Pump 3", "", "Exponential", 1e-3, None, None, "Vote"],
        ["Vote", "2oo3", None, None, None, None, "Fans"],
        ["Fans", "parallel", "Exponential", 2e-3, 2, None, "Backup"],
        ["Backup", "warm standby", "Exponential", 5e-4, None, None, "output"],
    ]})
    d = _one(data)
    nodes = _nodes(d.graph)
    assert nodes["Vote"]["type"] == "knode" and nodes["Vote"]["n"] == 2 and nodes["Vote"]["k"] == 3
    assert nodes["Fans"]["type"] == "parallel" and nodes["Fans"]["n"] == 2
    assert nodes["Backup"]["type"] == "standby" and nodes["Backup"]["dormancy"] == 0.5
    assert any("warm standby with no dormancy" in w for w in d.warnings)
    assert {("Supply", "Pump_1"), ("Supply", "Pump_3"), ("Backup", "output"), ("input", "Supply")} <= _edges(d.graph)
    normalize_graph(d.graph)


def test_fed_by_column_and_series_fallback():
    fed = _book({"Blocks": [
        ["ID", "Distribution", "alpha", "beta", "Fed by"],
        ["a", "Weibull", 100, 1.5, None],
        ["b", "Weibull", 100, 1.5, "a"],
    ]})
    assert _edges(_one(fed).graph) == {("input", "a"), ("a", "b"), ("b", "output")}

    plain = _book({"Blocks": [
        ["ID", "Distribution", "alpha", "beta"],
        ["a", "Weibull", 100, 1.5],
        ["b", "Weibull", 200, 1.5],
        ["c", "Weibull", 300, 1.5],
    ]})
    d = _one(plain)
    assert _edges(d.graph) == {("input", "a"), ("a", "b"), ("b", "c"), ("c", "output")}
    assert any("in series" in w for w in d.warnings)


def test_mu_sigma_and_other_distributions():
    data = _book({"Blocks": [
        ["ID", "Distribution", "mu", "sigma", "alpha", "beta"],
        ["n", "Normal", 500, 50, None, None],
        ["l", "log-normal", 6.2, 0.4, None, None],
        ["g", "Gamma", None, None, 3, 0.01],
        ["r", "Rayleigh", None, 300, None, None],
    ]})
    nodes = _nodes(_one(data).graph)
    assert nodes["n"]["model"]["distribution_id"] == "normal"
    assert nodes["l"]["model"]["params"] == [{"name": "mu", "value": 6.2}, {"name": "sigma", "value": 0.4}]
    assert nodes["g"]["model"]["params"] == [{"name": "alpha", "value": 3.0}, {"name": "beta", "value": 0.01}]
    assert nodes["r"]["model"]["params"] == [{"name": "sigma", "value": 300.0}]


# ---------------------------------------------------------------------------
# Repair
# ---------------------------------------------------------------------------


def test_repairable_when_every_block_has_a_repair_time():
    data = _book({"Blocks": [
        ["ID", "Type", "Distribution", "alpha", "beta", "MTTR", "Repair distribution", "Repair mu", "Repair sigma"],
        ["a", "", "Weibull", 1000, 1.5, 10, None, None, None],
        ["b", "", "Weibull", 1000, 1.5, None, "Lognormal", 2, 0.5],
        ["v", "1 of 2", None, None, None, None, None, None, None],
    ], "Connections": [["From", "To"], ["a", "v"], ["b", "v"]]})
    d = _one(data)
    assert d.graph["repairable"] is True
    nodes = _nodes(d.graph)
    assert nodes["a"]["repair"] == {"distribution_id": "exponential", "params": [{"name": "failure_rate", "value": 0.1}]}
    assert nodes["b"]["repair"]["distribution_id"] == "lognormal"
    assert "repair" not in nodes["v"]
    normalize_graph(d.graph)


def test_partial_repair_data_is_dropped_with_a_warning():
    data = _book({"Blocks": [
        ["ID", "Distribution", "alpha", "beta", "MTTR"],
        ["a", "Weibull", 1000, 1.5, 10],
        ["b", "Weibull", 1000, 1.5, None],
    ]})
    d = _one(data)
    assert not d.graph.get("repairable") and "repair" not in _nodes(d.graph)["a"]
    assert any("not every block has a repair time" in w and "“b”" in w for w in d.warnings)

    standby = _book({"Blocks": [
        ["ID", "Type", "Distribution", "failure_rate", "MTTR"],
        ["a", "", "Exponential", 1e-3, 10],
        ["s", "standby", "Exponential", 1e-3, 10],
    ]})
    d = _one(standby)
    assert not d.graph.get("repairable")
    assert any("doesn't support standby" in w for w in d.warnings)


# ---------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("sheets, message", [
    ({"Blocks": [["ID", "Distribution", "alpha", "beta"], ["a", "Weibull", 100, 1], ["a", "Weibull", 100, 1]]},
     "used twice"),
    ({"Blocks": [["ID", "Distribution", "alpha"], ["a", "Weibull", 100]]}, "needs alpha and beta"),
    ({"Blocks": [["ID", "Distribution", "alpha", "beta"], ["a", "Weibull", "lots", 1]]}, "isn't a number"),
    ({"Blocks": [["ID", "Distribution", "alpha", "beta"], ["a", "Weibull", -5, 1]]}, "greater than zero"),
    ({"Blocks": [["ID", "Distribution", "failure_rate"], ["a", "Pareto", 1]]}, "unsupported distribution"),
    ({"Blocks": [["ID", "Type", "Distribution", "failure_rate"], ["a", "gizmo", "Exponential", 1]]}, "unknown block type"),
    ({"Blocks": [["ID", "Type", "Distribution", "failure_rate"], ["a", "", "", None], ["b", "", "Exponential", 1]]},
     "no failure distribution"),
    ({"Blocks": [["ID", "failure_rate", "Feeds"], ["a", 1, "zzz"]]}, "isn't a block"),
    ({"Blocks": [["ID", "failure_rate", "Feeds"], ["a", 1, "b"], ["b", 1, "a"]]}, "loop"),
    ({"Blocks": [["ID", "failure_rate", "Feeds"], ["a", 1, "a"]]}, "connected to itself"),
    ({"Blocks": [["ID", "Type", "failure_rate", "Required", "Feeds"], ["a", "", 1, None, "v"], ["v", "k-of-n", None, 2, None]]},
     "needs 2 working branches"),
    ({"Blocks": [["ID", "Type", "Distribution", "failure_rate"], ["p", "parallel", "Exponential", 1]]}, "Count"),
])
def test_errors_are_user_facing(sheets, message):
    with pytest.raises(RbdImportError, match=message):
        _one(_book(sheets))


def test_unrecognised_layout_asks_for_a_mapping():
    data = _book({"Sheet1": [["Kit", "Life", "Shape"], ["Pump", 1000, 2]]})
    with pytest.raises(rx.ExcelNeedsMapping) as exc:
        _one(data)
    assert exc.value.code == "excel_mapping"


def test_explicit_mapping():
    data = _book({
        "Sheet1": [["Kit", "Life", "Shape", "Kind"], ["Pump", 1000, 2, "W"], ["Valve", 3000, 1, "W"]],
        "Wires": [["up", "down"], ["Pump", "Valve"]],
    })
    mapping = {
        "blocks": {"sheet": "Sheet1", "header_row": 1,
                   "mapping": {"label": "Kit", "alpha": "Life", "beta": "Shape"}},
        "connections": {"sheet": "Wires", "header_row": 1, "mapping": {"from": "up", "to": "down"}},
    }
    d = _one(data, mapping=mapping)
    nodes = _nodes(d.graph)
    assert nodes["Pump"]["model"]["distribution_id"] == "weibull"      # alpha + beta, no distribution column
    assert _edges(d.graph) == {("input", "Pump"), ("Pump", "Valve"), ("Valve", "output")}
    with pytest.raises(RbdImportError, match="ID or Name"):
        _one(data, mapping={"blocks": {"sheet": "Sheet1", "mapping": {"alpha": "Life"}}})
    with pytest.raises(RbdImportError, match="no column"):
        _one(data, mapping={"blocks": {"sheet": "Sheet1", "mapping": {"label": "Nope"}}})


def test_legacy_workbooks_are_refused_with_advice():
    with pytest.raises(RbdImportError, match="save it as .xlsx"):
        rbd_import.import_file(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "plant.xls")


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------

USER = {"uid": "user-xl", "email": "xl@example.org", "name": "XL"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def client(monkeypatch):
    from fastapi.testclient import TestClient

    from backend import config, db
    from backend.auth import get_current_user
    from backend.main import app

    monkeypatch.setattr(config, "AUTH_DISABLED", False)
    test_db = mongomock.MongoClient()["reliafy_test"]
    monkeypatch.setattr(db, "_db", test_db)
    monkeypatch.setattr(db, "_simulated", True)
    app.dependency_overrides[get_current_user] = lambda: USER
    tc = TestClient(app)
    tc.db = test_db
    try:
        yield tc
    finally:
        app.dependency_overrides.clear()


def test_template_endpoint(client):
    r = client.get("/api/rbds/import/template.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    assert "reliafy-rbd-template.xlsx" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content), read_only=True).sheetnames == ["README", "Blocks", "Connections"]


def test_import_endpoint_with_template(client):
    r = client.post("/api/rbds/import", files={"file": ("my.xlsx", rx.build_template(), XLSX)})
    assert r.status_code == 200, r.text
    (d,) = r.json()["diagrams"]
    assert d["source_format"] == "Excel" and d["n_nodes"] == 9
    assert client.db.rbds.count_documents({}) == 0


def test_import_endpoint_asks_for_and_accepts_a_mapping(client):
    data = _book({"Sheet1": [["Kit", "Life", "Shape"], ["Pump", 1000, 2]]})
    r = client.post("/api/rbds/import", files={"file": ("odd.xlsx", data, XLSX)})
    assert r.status_code == 422 and r.json()["code"] == "excel_mapping"
    mapping = {"blocks": {"sheet": "Sheet1", "header_row": 1,
                          "mapping": {"label": "Kit", "alpha": "Life", "beta": "Shape"}}}
    r = client.post("/api/rbds/import", files={"file": ("odd.xlsx", data, XLSX)}, data={"mapping": json.dumps(mapping)})
    assert r.status_code == 200, r.text
    assert r.json()["diagrams"][0]["n_nodes"] == 3
    bad = client.post("/api/rbds/import", files={"file": ("odd.xlsx", data, XLSX)}, data={"mapping": "{nope"})
    assert bad.status_code == 422
