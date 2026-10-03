"""Workbooks with cells far to the right or many merged ranges read quickly."""

import io
import re
import time
import zipfile



def _timed(fn, *args, **kwargs):
    start = time.perf_counter()
    try:
        out = fn(*args, **kwargs)
    except Exception as exc:  # noqa: BLE001 - the caller checks which
        out = exc
    return out, time.perf_counter() - start


def _with_merges(data: bytes, refs: list[str]) -> bytes:
    """``data`` with ``refs`` added as merged ranges on the first sheet."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                merges = "".join(f'<mergeCell ref="{r}"/>' for r in refs)
                body = re.sub(rb"<mergeCells[^>]*>.*?</mergeCells>", b"", body, flags=re.S)
                body = body.replace(b"</sheetData>",
                                    f'</sheetData><mergeCells count="{len(refs)}">{merges}</mergeCells>'.encode())
            dst.writestr(item, body)
    return out.getvalue()


def _sheet(rows: int, far_right: bool = False) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Time", "Status", "Group"])
    for r in range(2, rows + 2):
        ws.cell(r, 1, r)
        ws.cell(r, 2, 0)
        ws.cell(r, 3, "G" if r % 2 == 0 else None)
        if far_right:
            ws.cell(r, 16384, "x")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_workbook_with_far_right_cells_reads_quickly():
    from backend.services import excel

    data = _sheet(5000, far_right=True)
    table, seconds = _timed(excel.read_table, data, None, None, "wide.xlsx")
    assert not isinstance(table, Exception), table
    assert table.header == ["Time", "Status", "Group"]
    assert len(table.rows) == 5000
    assert any("first 200 columns" in n for n in table.notes)
    assert seconds < 3.0
    info = excel.inspect(data, "wide.xlsx")
    assert info["sheets"][0]["cols"] == 3


def test_workbook_merged_ranges_fill_down_within_the_cap():
    from backend.services import excel

    data = _with_merges(_sheet(20), [f"C{r}:C{r + 1}" for r in range(2, 20, 2)])
    table = excel.read_table(data, None, None, "m.xlsx")
    assert [r[2] for r in table.rows[:4]] == ["G", "G", "G", "G"]
    assert not any("merged" in n for n in table.notes)


def test_workbook_with_many_merged_ranges_reads_quickly_without_merge_fill(monkeypatch):
    from backend.services import excel

    monkeypatch.setattr(excel, "MAX_MERGES", 1000)
    # Overlapping ranges over every row, plus some below the rows read.
    refs = [f"C2:C{2000 + i}" for i in range(1500)] + [f"A{excel.MAX_ROWS + 5}:A{excel.MAX_ROWS + 9}"]
    data = _with_merges(_sheet(2000), refs)
    table, seconds = _timed(excel.read_table, data, None, None, "m.xlsx")
    assert not isinstance(table, Exception), table
    assert any("merged cell ranges" in n for n in table.notes)
    assert seconds < 3.0


def test_overlapping_merged_ranges_fill_one_value_per_column():
    from backend.services import excel

    refs = [f"C2:C{2000 + i}" for i in range(500)]
    data = _with_merges(_sheet(2000), refs)
    table, seconds = _timed(excel.read_table, data, None, None, "m.xlsx")
    assert not isinstance(table, Exception), table
    # The first range (C2:C2000) fills; the overlapping ones are ignored.
    assert {r[2] for r, n in zip(table.rows, table.row_numbers) if n <= 2000} == {"G"}
    assert seconds < 3.0
